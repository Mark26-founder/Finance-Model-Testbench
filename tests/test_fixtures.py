"""
Tests for C2 — Financial Test Fixtures Integrity and Manifest Verification
"""

import json
import os
import openpyxl


def test_fixture_directory_and_manifest_exist():
    """Verify all fixture files and manifests exist in expected paths."""
    manifest_path = os.path.join("examples", "fixtures", "manifests", "ground_truth_manifest.json")
    assert os.path.exists(manifest_path)
    
    with open(manifest_path, "r", encoding="utf-8") as f:
        manifest = json.load(f)
        
    assert "fixtures" in manifest
    assert len(manifest["fixtures"]) == 5

    for item in manifest["fixtures"]:
        file_path = item["relative_path"]
        assert os.path.exists(file_path), f"Fixture file missing: {file_path}"


def test_valid_fixture_structure():
    """Verify valid fixture contains expected sheets and coherent formulas."""
    file_path = os.path.join("examples", "fixtures", "valid", "valid_three_statement.xlsx")
    wb = openpyxl.load_workbook(file_path, data_only=False)
    
    expected_sheets = ["Assumptions", "Income Statement", "Balance Sheet", "Cash Flow Statement"]
    for sheet in expected_sheets:
        assert sheet in wb.sheetnames

    ws_bs = wb["Balance Sheet"]
    # Check Balance Sheet Total Assets formula
    assert ws_bs["C5"].value == "=SUM(C2:C4)"
    # Check Balance Sheet Retained Earnings formula
    assert ws_bs["C10"].value == "=B10+'Income Statement'!C10"


def test_defective_bs_imbalance_fixture():
    """Verify defective BS imbalance fixture contains expected defect formula."""
    file_path = os.path.join("examples", "fixtures", "defective", "defective_balance_sheet_imbalance.xlsx")
    wb = openpyxl.load_workbook(file_path, data_only=False)
    ws_bs = wb["Balance Sheet"]
    assert ws_bs["C5"].value == "=SUM(C2:C4)+50.0"


def test_defective_re_rollforward_fixture():
    """Verify defective Retained Earnings fixture contains expected defect formula."""
    file_path = os.path.join("examples", "fixtures", "defective", "defective_retained_earnings_rollforward.xlsx")
    wb = openpyxl.load_workbook(file_path, data_only=False)
    ws_bs = wb["Balance Sheet"]
    assert ws_bs["C10"].value == "=B10+'Income Statement'!C10-100.0"


def test_defective_cash_flow_linkage_fixture():
    """Verify defective Cash Flow Linkage fixture contains hardcoded 0.0 defect."""
    file_path = os.path.join("examples", "fixtures", "defective", "defective_cash_flow_linkage.xlsx")
    wb = openpyxl.load_workbook(file_path, data_only=False)
    ws_cf = wb["Cash Flow Statement"]
    assert ws_cf["B3"].value == 0.0
